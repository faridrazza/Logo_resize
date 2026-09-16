"""Pull the logos embedded in an xlsx into data/input.

    python scripts/fetch_sheet_images.py logo_resize_test_100.xlsx

Unlike fetch_sheet.py (which reads a column of URLs), this reads pictures anchored
in the worksheet, names each one by its row number and brand name, and writes a
manifest so the report can label rows the way the sheet does.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl                    # noqa: E402

from config import settings        # noqa: E402

PREVIEW_COL = 1                    # column B, "Original Preview"


def slug(text: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", str(text or "")).strip("-").lower()
    return s or "logo"


def main() -> None:
    xlsx = Path(sys.argv[1] if len(sys.argv) > 1 else "logo_resize_test_100.xlsx")
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[wb.sheetnames[0]]

    # row number -> (brand, category, stated width, stated height)
    meta: dict[int, tuple] = {}
    for row in ws.iter_rows(min_row=2, values_only=False):
        cells = [c.value for c in row]
        if not cells or cells[0] in (None, ""):
            continue
        try:
            no = int(str(cells[0]).strip())
        except ValueError:
            continue
        meta[no] = (cells[2], cells[3], cells[4], cells[5])

    images = getattr(ws, "_images", [])
    print(f"{xlsx.name}: {len(meta)} data rows, {len(images)} embedded pictures\n")

    by_row: dict[int, list] = {}
    for img in images:
        frm = getattr(img.anchor, "_from", None)
        if frm is None or frm.col != PREVIEW_COL:
            continue
        by_row.setdefault(frm.row + 1, []).append(img)   # anchor rows are 0-based

    skipped = [r for r, v in by_row.items() if len(v) > 1]
    lines, written = [], 0
    for no in sorted(meta):
        sheet_row = no + 1                                # row 1 is the header
        found = by_row.get(sheet_row)
        if not found:
            print(f"  [{no:>3}] no picture anchored in row {sheet_row}")
            lines.append(f"{no}\tMISSING\t{meta[no][0]}\t{meta[no][1]}\t{meta[no][2]}\t{meta[no][3]}")
            continue
        img = found[0]
        data = img.ref.getvalue() if hasattr(img.ref, "getvalue") else Path(img.ref).read_bytes()
        name = f"{no:03d}_{slug(meta[no][0])}.{(img.format or 'png').lower()}"
        (settings.INPUT_DIR / name).write_bytes(data)
        written += 1
        brand, cat, w, h = meta[no]
        lines.append(f"{no}\t{name}\t{brand}\t{cat}\t{w}\t{h}")
        if no <= 5 or no % 20 == 0:
            print(f"  [{no:>3}] {name:<40} {len(data):>8,} B   {brand} / {cat}")

    if skipped:
        print(f"\n  rows with more than one picture (first used): {skipped}")

    manifest = settings.INPUT_DIR / "_manifest.tsv"
    manifest.write_text("\n".join(lines), encoding="utf-8")
    print(f"\n{written} logos written to {settings.INPUT_DIR}")
    print(f"Manifest -> {manifest}")


if __name__ == "__main__":
    main()
